import csv
from io import BytesIO, StringIO
from pathlib import Path
from uuid import UUID

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.ext.asyncio import AsyncSession

from aiimage.assets.models import Asset
from aiimage.bulk.models import BulkJob, BulkJobRow
from aiimage.bulk.schemas import BulkJobResponse, BulkJobRowResponse
from aiimage.catalog.models import (
    Product,
    ProductCategory,
    ProductReference,
    ReferenceView,
    TruthAnchor,
)
from aiimage.models.domain import Capability
from aiimage.models.models import ModelConfiguration
from aiimage.templates.models import PackKind, TemplatePack, TemplatePackVersion
from aiimage.templates.schemas import CompilePlanRequest
from aiimage.templates.service import create_production_plan, ensure_first_party_packs
from aiimage.workflow.queue import QueueHints
from aiimage.workflow.service import execute_production_plan

REQUIRED_COLUMNS = {"sku", "name", "category", "platform_slug"}
MAX_CSV_BYTES = 2 * 1024 * 1024
MAX_ROWS = 500


class BulkImportError(ValueError):
    pass


def parse_csv(content: bytes) -> list[dict[str, str]]:
    if len(content) > MAX_CSV_BYTES:
        raise BulkImportError("CSV exceeds 2 MiB")
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise BulkImportError("CSV must be UTF-8") from exc
    reader = csv.DictReader(StringIO(text))
    columns = set(reader.fieldnames or [])
    missing = REQUIRED_COLUMNS - columns
    if missing:
        raise BulkImportError(f"Missing CSV columns: {', '.join(sorted(missing))}")
    rows = [{key: (value or "").strip() for key, value in row.items()} for row in reader]
    if not rows:
        raise BulkImportError("CSV contains no data rows")
    if len(rows) > MAX_ROWS:
        raise BulkImportError(f"CSV exceeds {MAX_ROWS} rows")
    return rows


def parse_xlsx(content: bytes) -> list[dict[str, str]]:
    if len(content) > MAX_CSV_BYTES:
        raise BulkImportError("XLSX exceeds 2 MiB")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(values)]
    except (OSError, StopIteration, ValueError) as exc:
        raise BulkImportError("XLSX is empty or invalid") from exc
    missing = REQUIRED_COLUMNS - set(headers)
    if missing:
        raise BulkImportError(f"Missing XLSX columns: {', '.join(sorted(missing))}")
    rows = [
        {
            header: "" if value is None else str(value).strip()
            for header, value in zip(headers, values_row, strict=False)
            if header
        }
        for values_row in values
        if any(value is not None and str(value).strip() for value in values_row)
    ]
    if not rows:
        raise BulkImportError("XLSX contains no data rows")
    if len(rows) > MAX_ROWS:
        raise BulkImportError(f"XLSX exceeds {MAX_ROWS} rows")
    return rows


def parse_import(filename: str, content: bytes) -> list[dict[str, str]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        return parse_csv(content)
    if suffix == ".xlsx":
        return parse_xlsx(content)
    raise BulkImportError("Only .csv and .xlsx files are supported")


async def _pack_version_by_slug(
    session: AsyncSession, slug: str, kind: PackKind
) -> TemplatePackVersion | None:
    return (
        await session.execute(
            select(TemplatePackVersion)
            .join(TemplatePack, TemplatePack.id == TemplatePackVersion.pack_id)
            .where(
                TemplatePack.slug == slug,
                TemplatePack.kind == kind.value,
                TemplatePackVersion.status == "published",
            )
            .order_by(TemplatePackVersion.version.desc())
            .limit(1)
        )
    ).scalar_one_or_none()


async def _validate_row(
    session: AsyncSession, row: dict[str, str]
) -> tuple[Product | None, TemplatePackVersion, Asset | None, ProductCategory, ReferenceView]:
    sku = row["sku"].strip().upper()
    if not sku or not row["name"]:
        raise BulkImportError("SKU and name are required")
    try:
        category = ProductCategory(row["category"])
        view = ReferenceView(row.get("reference_view") or "front")
    except ValueError as exc:
        raise BulkImportError("Unsupported category or reference_view") from exc
    platform_version = await _pack_version_by_slug(
        session, row["platform_slug"], PackKind.PLATFORM
    )
    if platform_version is None:
        raise BulkImportError("Unknown published platform_slug")
    product = await session.scalar(select(Product).where(Product.sku == sku))
    asset = None
    raw_asset_id = row.get("reference_asset_id", "")
    if raw_asset_id:
        try:
            asset = await session.get(Asset, UUID(raw_asset_id))
        except ValueError as exc:
            raise BulkImportError("reference_asset_id is not a UUID") from exc
        if asset is None or not asset.mime_type.startswith("image/"):
            raise BulkImportError("reference_asset_id is not an image asset")
    has_reference = False
    if product is not None:
        has_reference = (
            await session.scalar(
                select(ProductReference.id).where(ProductReference.product_id == product.id).limit(1)
            )
        ) is not None
    if not has_reference and asset is None:
        raise BulkImportError("New products and products without images require reference_asset_id")
    return product, platform_version, asset, category, view


def _to_response(job: BulkJob, rows: list[BulkJobRow]) -> BulkJobResponse:
    return BulkJobResponse(
        id=job.id,
        filename=job.filename,
        dry_run=job.dry_run,
        status=job.status,
        model_configuration_id=job.model_configuration_id,
        category_pack_version_id=job.category_pack_version_id,
        brand_pack_version_id=job.brand_pack_version_id,
        total_rows=job.total_rows,
        succeeded_rows=job.succeeded_rows,
        failed_rows=job.failed_rows,
        created_at=job.created_at,
        rows=[BulkJobRowResponse.model_validate(row) for row in rows],
    )


async def create_bulk_job(
    session: AsyncSession,
    queue: QueueHints,
    *,
    filename: str,
    content: bytes,
    dry_run: bool,
    model_configuration_id: UUID,
    category_pack_version_id: UUID,
    brand_pack_version_id: UUID,
    user_id: UUID,
) -> BulkJobResponse:
    parsed_rows = parse_import(filename, content)
    await ensure_first_party_packs(session)
    model = await session.get(ModelConfiguration, model_configuration_id)
    category_pack = await session.get(TemplatePackVersion, category_pack_version_id)
    brand_pack = await session.get(TemplatePackVersion, brand_pack_version_id)
    if (
        model is None
        or not model.is_enabled
        or Capability.REFERENCE_TO_IMAGE.value not in model.capabilities
    ):
        raise BulkImportError("Enabled reference_to_image model is required")
    if category_pack is None or brand_pack is None:
        raise BulkImportError("Category or brand pack version not found")
    kinds = dict(
        (
            await session.execute(
                select(TemplatePackVersion.id, TemplatePack.kind)
                .join(TemplatePack, TemplatePack.id == TemplatePackVersion.pack_id)
                .where(
                    TemplatePackVersion.id.in_(
                        [category_pack_version_id, brand_pack_version_id]
                    )
                )
            )
        ).all()
    )
    if kinds.get(category_pack_version_id) != PackKind.CATEGORY.value:
        raise BulkImportError("category_pack_version_id is not a category pack")
    if kinds.get(brand_pack_version_id) != PackKind.BRAND.value:
        raise BulkImportError("brand_pack_version_id is not a brand pack")

    job = BulkJob(
        filename=filename[:255],
        dry_run=dry_run,
        status="validating" if dry_run else "processing",
        model_configuration_id=model_configuration_id,
        category_pack_version_id=category_pack_version_id,
        brand_pack_version_id=brand_pack_version_id,
        total_rows=len(parsed_rows),
        created_by_user_id=user_id,
    )
    session.add(job)
    await session.flush()
    job_id = job.id
    await session.commit()
    await session.refresh(job)
    result_rows: list[BulkJobRow] = []
    for index, row in enumerate(parsed_rows, start=2):
        result = BulkJobRow(
            job_id=job_id,
            row_number=index,
            sku=row.get("sku", "").upper(),
            status="validating",
            input_data=row,
            batch_ids=[],
        )
        session.add(result)
        result_rows.append(result)
        try:
            product, platform_version, asset, category, view = await _validate_row(session, row)
            if dry_run:
                result.status = "valid"
                continue
            if product is None:
                product = Product(
                    sku=row["sku"].upper(),
                    name=row["name"],
                    category=category.value,
                    brand=row.get("brand") or None,
                    created_by_user_id=user_id,
                )
                session.add(product)
                await session.flush()
                session.add(
                    TruthAnchor(
                        product_id=product.id,
                        version=1,
                        document={"source": "bulk_import", "bulk_job_id": str(job_id)},
                        created_by_user_id=user_id,
                    )
                )
            if asset is not None:
                existing_reference = await session.scalar(
                    select(ProductReference.id).where(
                        ProductReference.product_id == product.id,
                        ProductReference.asset_id == asset.id,
                        ProductReference.view == view.value,
                    )
                )
                if existing_reference is None:
                    session.add(
                        ProductReference(product_id=product.id, asset_id=asset.id, view=view.value)
                    )
            await session.commit()
            plan = await create_production_plan(
                session,
                payload=CompilePlanRequest(
                    product_id=product.id,
                    platform_pack_version_id=platform_version.id,
                    category_pack_version_id=category_pack_version_id,
                    brand_pack_version_id=brand_pack_version_id,
                    mode=row.get("mode") or "strict",
                ),
                user_id=user_id,
            )
            batches = await execute_production_plan(
                session,
                queue,
                plan_id=plan.id,
                model_configuration_id=model_configuration_id,
                user_id=user_id,
            )
            result.product_id = product.id
            result.production_plan_id = plan.id
            result.batch_ids = [str(batch.id) for batch in batches]
            result.status = "queued"
        except (BulkImportError, ValueError, RuntimeError, SQLAlchemyError) as exc:
            await session.rollback()
            result = await session.merge(result)
            result.status = "invalid" if dry_run else "failed"
            result.error = str(exc)[:1000]
        await session.commit()

    job = await session.get(BulkJob, job_id)
    assert job is not None
    rows = list(
        (
            await session.scalars(
                select(BulkJobRow)
                .where(BulkJobRow.job_id == job.id)
                .order_by(BulkJobRow.row_number)
            )
        ).all()
    )
    job.succeeded_rows = sum(row.status in {"valid", "queued"} for row in rows)
    job.failed_rows = len(rows) - job.succeeded_rows
    job.status = (
        "validated"
        if dry_run and not job.failed_rows
        else "validation_failed"
        if dry_run
        else "queued"
        if not job.failed_rows
        else "partially_queued"
        if job.succeeded_rows
        else "failed"
    )
    await session.commit()
    return _to_response(job, rows)


async def get_bulk_job(session: AsyncSession, job_id: UUID) -> BulkJobResponse | None:
    job = await session.get(BulkJob, job_id)
    if job is None:
        return None
    rows = list(
        (
            await session.scalars(
                select(BulkJobRow)
                .where(BulkJobRow.job_id == job.id)
                .order_by(BulkJobRow.row_number)
            )
        ).all()
    )
    return _to_response(job, rows)


async def list_bulk_jobs(session: AsyncSession) -> list[BulkJobResponse]:
    jobs = list((await session.scalars(select(BulkJob).order_by(BulkJob.created_at.desc()))).all())
    return [response for job in jobs if (response := await get_bulk_job(session, job.id))]
